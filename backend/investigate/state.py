"""Deterministic core of investigation mode.

The agent (an LLM) decides what to look at and proposes conclusions. This module
owns everything that must be trustworthy: it runs the verified pipeline, checks
every quote against the source posts, detects contradictions, refuses to record
a 'shared agenda' while contradictions are open, and renders the report."""

from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path

from app.verification import quote_in_text

MAX_POSTS = 10


@dataclass
class Finding:
    theme: str
    evidence: dict[str, str]  # post_id -> verified verbatim quote


@dataclass
class Contradiction:
    post_ids: list[str]
    description: str
    source: str  # "code" (detected deterministically) or "agent"


@dataclass
class Investigation:
    posts: dict[str, str]
    runner: object  # async callable: text -> PipelineResult
    results: dict = field(default_factory=dict)
    findings: list[Finding] = field(default_factory=list)
    contradictions: list[Contradiction] = field(default_factory=list)
    rejected: list[str] = field(default_factory=list)
    summary: str = ""
    cost_usd: float = 0.0

    def __post_init__(self):
        if not 1 < len(self.posts) <= MAX_POSTS:
            raise ValueError(f"investigation needs 2..{MAX_POSTS} posts")

    @property
    def halted(self) -> bool:
        return bool(self.contradictions)

    async def analyze_post(self, post_id: str):
        if post_id not in self.posts:
            raise KeyError(f"unknown post_id {post_id!r}; known: {sorted(self.posts)}")
        if post_id not in self.results:
            res = await self.runner(self.posts[post_id])
            self.results[post_id] = res
            self.cost_usd += res.usage.cost_usd
            self._detect_contradictions()
        return self.results[post_id]

    def _detect_contradictions(self):
        """Code-level check the agent cannot talk its way around: posts that the
        pipeline classifies in opposite directions cannot be pooled into one narrative."""
        verdicts = {pid: r.verdict for pid, r in self.results.items()}
        ai = [p for p, v in verdicts.items() if v == "ai"]
        human = [p for p, v in verdicts.items() if v == "human"]
        if ai and human and not any(c.source == "code" for c in self.contradictions):
            self.contradictions.append(
                Contradiction(
                    sorted(ai + human),
                    f"Verdicts disagree: AI-like {sorted(ai)} vs human-like {sorted(human)}.",
                    "code",
                )
            )

    def flag_contradiction(self, post_ids: list[str], description: str):
        unknown = [p for p in post_ids if p not in self.posts]
        if unknown:
            raise KeyError(f"unknown post_ids {unknown}")
        self.contradictions.append(Contradiction(post_ids, description, "agent"))

    def record_finding(self, theme: str, evidence: dict[str, str]) -> tuple[bool, str]:
        """A shared-agenda finding must (a) not be made while contradictions are open,
        (b) span >= 2 posts, (c) quote every post, with each quote verbatim in that post."""
        if self.halted:
            return False, "REFUSED: open contradictions. Stop synthesising; call finish() so a human can review."
        if len(evidence) < 2:
            return False, "REFUSED: a shared agenda needs evidence from at least 2 posts."
        bad = [pid for pid, q in evidence.items() if pid not in self.posts or not quote_in_text(q, self.posts[pid])]
        if bad:
            self.rejected.append(theme)
            return False, f"REFUSED: quote not found verbatim in post(s) {bad}. Finding discarded."
        self.findings.append(Finding(theme, evidence))
        return True, "recorded"

    @property
    def status(self) -> str:
        if self.halted:
            return "ESCALATED: contradictions found, human review required"
        if len(self.results) < len(self.posts):
            return "INCOMPLETE: not all posts were analysed"
        return "COMPLETE"

    def write_docx(self, path: Path):
        from docx import Document

        doc = Document()
        doc.add_heading("Deslopify investigation report", 0)
        doc.add_paragraph(f"Generated {datetime.now(UTC):%Y-%m-%d %H:%M} UTC. Status: {self.status}")
        doc.add_paragraph(
            "AI-text detection is unreliable. Nothing here is proof of authorship or intent; " "treat it as a lead for human review."
        ).runs[0].italic = True

        if self.halted:
            doc.add_heading("Contradictions (analysis stopped)", 1)
            for c in self.contradictions:
                doc.add_paragraph(f"[{c.source}] {', '.join(c.post_ids)}: {c.description}", style="List Bullet")

        doc.add_heading("Per-post results", 1)
        table = doc.add_table(rows=1, cols=4)
        table.style = "Light Grid Accent 1"
        for cell, h in zip(table.rows[0].cells, ["Post", "Verdict", "P(AI)", "Verified primary goal"], strict=True):
            cell.text = h
        for pid in sorted(self.posts):
            r = self.results.get(pid)
            cells = table.add_row().cells
            cells[0].text = pid
            if r is None:
                cells[1].text = "not analysed"
                continue
            cells[1].text, cells[2].text = r.verdict, f"{r.ai_probability:.2f}"
            cells[3].text = r.agenda.primary_goal.claim if r.agenda else f"n/a ({r.agenda_status})"

        doc.add_heading("Shared agenda findings (quotes verified by code)", 1)
        if not self.findings:
            doc.add_paragraph("None recorded." + (" Analysis halted before synthesis." if self.halted else ""))
        for f in self.findings:
            doc.add_paragraph(f.theme, style="List Bullet")
            for pid, q in f.evidence.items():
                doc.add_paragraph(f"{pid}: “{q}”", style="List Bullet 2")
        if self.rejected:
            doc.add_paragraph(f"{len(self.rejected)} proposed finding(s) discarded for unverifiable quotes.")

        if self.summary:
            doc.add_heading("Agent narrative (unverified, written by the model)", 1)
            doc.add_paragraph(self.summary)
        doc.add_paragraph(f"Estimated API cost of pipeline calls: ${self.cost_usd:.4f}")
        doc.save(path)

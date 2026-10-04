"""Weekly usage report: /stats -> pandas -> xlsx + docx -> email.

    python -m reports.weekly_report            # fetch, build, email
    python -m reports.weekly_report --no-email # just write files to ./out

Env: STATS_URL, STATS_SECRET (sent as a header), SMTP_HOST, SMTP_PORT (587),
SMTP_USER, SMTP_PASSWORD, REPORT_FROM, REPORT_TO.
Scheduled by .github/workflows/weekly-report.yml.
"""

import argparse
import os
import smtplib
from email.message import EmailMessage
from pathlib import Path

import pandas as pd
import requests
from docx import Document


def fetch_stats(url: str, secret: str, days: int = 14) -> dict:
    r = requests.get(f"{url.rstrip('/')}/stats", params={"days": days}, headers={"X-Stats-Secret": secret}, timeout=30)
    r.raise_for_status()
    return r.json()


def build_frames(stats: dict) -> tuple[pd.DataFrame, dict]:
    """Returns (daily frame sorted by date with 7-day rolling mean, summary dict).
    The last 7 days are compared with the 7 before them."""
    df = pd.DataFrame(stats["daily"]).sort_values("date").reset_index(drop=True)
    df["analyses_7d_avg"] = df["analyses"].rolling(7, min_periods=1).mean().round(1)
    this_week, prev_week = df.tail(7), df.iloc[-14:-7]

    def pct_change(a: float, b: float):
        return None if b == 0 else round((a - b) / b * 100, 1)

    summary = {
        "analyses_this_week": int(this_week["analyses"].sum()),
        "analyses_prev_week": int(prev_week["analyses"].sum()),
        "wow_change_pct": pct_change(this_week["analyses"].sum(), prev_week["analyses"].sum()),
        "cost_usd_this_week": round(float(this_week["cost_usd"].sum()), 2),
        "errors_this_week": int(this_week["errors"].sum()),
        "error_rate_pct": round(float(this_week["errors"].sum()) / max(int(this_week["analyses"].sum()), 1) * 100, 1),
        "avg_cost_per_analysis_usd": round(float(this_week["cost_usd"].sum()) / max(int(this_week["analyses"].sum()), 1), 4),
        "all_time_total": stats["total_analyses"],
        "all_time_ai_pct": stats["ai_detection_rate_pct"],
        "all_time_uncertain": stats.get("uncertain", 0),
    }
    return df, summary


def write_xlsx(df: pd.DataFrame, summary: dict, path: Path):
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        pd.DataFrame(list(summary.items()), columns=["metric", "value"]).to_excel(xw, sheet_name="Summary", index=False)
        df.to_excel(xw, sheet_name="Daily", index=False)
        for ws in xw.book.worksheets:
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = max(len(str(c.value or "")) for c in col) + 3


def write_docx(df: pd.DataFrame, summary: dict, path: Path):
    doc = Document()
    doc.add_heading("Deslopify: weekly report", 0)
    doc.add_paragraph(f"Period: {df['date'].iloc[-7]} to {df['date'].iloc[-1]}")
    doc.add_heading("Summary", 1)
    for k, v in summary.items():
        doc.add_paragraph(f"{k.replace('_', ' ')}: {v}", style="List Bullet")
    doc.add_heading("Daily", 1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    for cell, h in zip(table.rows[0].cells, ["Date", "Analyses", "Cost (USD)", "Errors"], strict=True):
        cell.text = h
    for _, r in df.tail(7).iterrows():
        cells = table.add_row().cells
        cells[0].text, cells[1].text = str(r["date"]), str(int(r["analyses"]))
        cells[2].text, cells[3].text = f"{r['cost_usd']:.2f}", str(int(r["errors"]))
    doc.save(path)


def send_email(files: list[Path], summary: dict):
    msg = EmailMessage()
    msg["Subject"] = f"Deslopify weekly report: {summary['analyses_this_week']} analyses"
    msg["From"], msg["To"] = os.environ["REPORT_FROM"], os.environ["REPORT_TO"]
    msg.set_content("\n".join(f"{k}: {v}" for k, v in summary.items()))
    for f in files:
        msg.add_attachment(f.read_bytes(), maintype="application", subtype="octet-stream", filename=f.name)
    with smtplib.SMTP(os.environ["SMTP_HOST"], int(os.getenv("SMTP_PORT", "587"))) as s:
        s.starttls()
        s.login(os.environ["SMTP_USER"], os.environ["SMTP_PASSWORD"])
        s.send_message(msg)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--no-email", action="store_true")
    ap.add_argument("--out", default="out")
    a = ap.parse_args()

    stats = fetch_stats(os.environ["STATS_URL"], os.environ["STATS_SECRET"])
    df, summary = build_frames(stats)
    out = Path(a.out)
    out.mkdir(exist_ok=True)
    files = [out / "weekly_report.xlsx", out / "weekly_report.docx"]
    write_xlsx(df, summary, files[0])
    write_docx(df, summary, files[1])
    if not a.no_email:
        send_email(files, summary)
    print("report written" + ("" if a.no_email else " and emailed"), summary)


if __name__ == "__main__":
    main()

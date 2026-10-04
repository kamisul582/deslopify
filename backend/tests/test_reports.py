import time

from docx import Document
from openpyxl import load_workbook

from reports.weekly_report import build_frames, write_docx, write_xlsx


def fake_stats():
    daily = []
    for i in range(14):
        d = time.strftime("%Y-%m-%d", time.gmtime(time.time() - i * 86400))
        daily.append({"date": d, "analyses": 10 if i < 7 else 5, "cost_usd": 0.02 * (10 if i < 7 else 5), "errors": 1 if i == 0 else 0})
    return {"total_analyses": 105, "ai_detected": 40, "not_ai": 60, "uncertain": 5, "ai_detection_rate_pct": 38.1, "daily": daily}


def test_summary_and_files(tmp_path):
    df, s = build_frames(fake_stats())
    assert s["analyses_this_week"] == 70 and s["analyses_prev_week"] == 35 and s["wow_change_pct"] == 100.0
    assert s["errors_this_week"] == 1 and s["cost_usd_this_week"] == 1.4
    write_xlsx(df, s, tmp_path / "r.xlsx")
    write_docx(df, s, tmp_path / "r.docx")
    wb = load_workbook(tmp_path / "r.xlsx")
    assert wb.sheetnames == ["Summary", "Daily"] and wb["Daily"].max_row == 15
    assert "weekly report" in Document(tmp_path / "r.docx").paragraphs[0].text.lower()
